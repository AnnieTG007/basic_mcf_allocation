"""把仿真状态和已算好的统计写成资源回放 JSON、Excel 工作簿与图表。

由 traffic_scan 调用；不生成业务、不决定分配、不推进时间。输出路径由调用方给出。
回放 kind=qkd_resource_timeline、schema_version=2，config 说明频率、芯分组、
功率 dBm 和仿真时间单位。initial_occupancy 为初始状态，states 为每次资源
改变后的 {time, event, group_id, occupancy}，包含预热段。终点清空事件为
horizon_release，带 group_ids，仅结束回放，不冒充自然离去、不改变仿真统计。

本回放用于 --export-business 全网三芯实验，仅保留经过所选链路的已接入业务。
完整路由保留在 path，hops 和 occupancy 只含所选链路；全网阻塞统计在 metrics。
business_groups 每条记录是一组到达：
group_id、source/destination、direction、arrival_time、holding_time、scheduled_end_time、
status（仅 accepted）、release_time、release_reason（natural/horizon）。
source/destination/direction 描述端到端业务；回放传播方向必须读取 hops.direction，
不能由端到端节点编号推断。allocation 含完整 path 和选中链路的 hops（link/direction/cores/physical_cores）、
channel_index、frequency_hz 和 wavelength_nm。cores 为零起始仿真编号，physical_cores
为物理芯号，不包含硬件端口。同组同方向三芯占用一致，两个方向可复用同一波长。
负载单位为三芯业务组 Erlang，到达率按组/仿真时间计，功率为每芯每信道输入功率。

occupancy 的维度为 [有向链路, 仿真芯, 经典信道]，链路顺序见 directed_links。
-2 表示该方向不可用，-1 表示空闲，非负整数为占用业务 ID（0 也表示占用）。
量子资源只写在 config 中；经典索引从 0 开始，长度由经典信道数决定，
默认 10 个依次对应 C40 至 C36、C34、C32 至 C29（跳过 C33），
不是 ITU 编号，也不是包含量子频率的仿真内部索引。channel_spacing_hz 只是
基础网格间隔；跨量子频段及跳过 C33 均形成缺口，应读取实际 classical_frequencies_hz。

扫描工作簿仅含 Summary；业务工作簿仅含 LoadSweep/PowerSweep。表格为普通黑白
单元格，仅含条件、算法、SKR/OSNR/阻塞率/协同度及四项各自相对FF的比值。
完整配置、逐种子数据与样本保留在JSON。JSON 的 SKR 为 bit/s，趋势图和表格
为 kbit/s；它们表示密钥生成速率而非累计密钥比特。误差范围是种子
均值间的样本标准差，单种子时不显示；空白表示未定义而非零。
"""
from copy import deepcopy
from hashlib import sha256
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd

from skr_calculation import skr_model_config


TRACE_SCHEMA_VERSION = 2
# 七芯专用映射：仿真 0..5 是外圈物理 2..7，仿真 6 是物理中心芯 1。
CORE_TO_PHYSICAL = [2, 3, 4, 5, 6, 7, 1]


class TrafficRecorder:
    """按事件发生后的资源归属记录回放，不参与分配。

    occupancy[有向链路][仿真芯][经典信道] 保存业务 ID，阻塞事件不改变占用。
    固定量子芯/频率单独放在 config 中。物理芯映射只适用于七芯布局。
    """

    def __init__(self, sim, *, seed, warmup):
        if not sim.allocator.bind_three:
            raise ValueError('Business trace requires a three-core experiment allocator')
        self.sim = sim
        a, b = sim.observed_link
        self.links = [(a, b), (b, a)]
        self.link_index = {tuple(link): i for i, link in enumerate(self.links)}
        self.occupancy = [
            [[-1 if sim.m_resourceMap[a, b, c, w] == 1 else -2
              for w in range(sim.quantum_wave_num, sim.WaveNumber)]
             for c in range(sim.core_num)] for a, b in self.links]
        self.config = dict(
            algorithm='first-fit' if sim.algorithm == 'FF' else sim.algorithm.lower(),
            seed=int(seed), duration=float(sim.Ts), warmup=float(warmup),
            observed_link=list(sim.observed_link), allocation_scope='full_network',
            network_node_count=len(sim.graph), network_edge_count=sim.graph.number_of_edges(),
            network_length_scaling=deepcopy(sim.graph.graph['length_scaling']),
            network_edges_m=[(int(a), int(b), float(sim.a_m[a, b])) for a, b in sim.graph.edges],
            mean_holding_time=float(sim.m_rou1), arrival_rate=float(sim.lambda1),
            offered_load_erlang=float(sim.lambda1 * sim.m_rou1),
            power_dbm=float(10 * math.log10(sim.launch_power / 1e-3)),
            power_reference='per_core_per_classical_channel_at_fiber_input',
            allocation_mode='three_core_bound', traffic_unit='three_core_business_group',
            load_unit='Erlang of three-core business groups',
            arrival_rate_unit='three-core groups per simulation unit',
            direction_definition='business direction uses end-to-end source/destination; each hop direction uses its link endpoints (forward u<v, backward u>v); replay must use hop direction',
            time_unit='simulation_unit', direction_mode='bidirectional',
            directed_links=[list(link) for link in self.links],
            link_lengths_m=[float(sim.a_m[a, b]) for a, b in self.links],
            core_to_physical=CORE_TO_PHYSICAL,
            forward_cores=list(sim.classical_forward_cores),
            backward_cores=list(sim.classical_backward_cores),
            quantum_cores=list(sim.quantum_cores),
            quantum_frequencies_hz=[int(f) for f in sim.available_channel[:sim.quantum_wave_num]],
            classical_frequencies_hz=[int(f) for f in sim.available_channel[sim.quantum_wave_num:]],
            channel_spacing_hz=int(sim.wave_interval),
            end_policy='release_at_horizon',
            skr_model=skr_model_config(sim.bb84_params, sim.detector_params),
        )
        self.initial = deepcopy(self.occupancy)
        self.active = {}
        self.states = []
        self.business_groups = []

    def record(self, event, *, blocked=False):
        """记录经过观测链路的已接入组及其真实资源变化；阻塞和其他链路业务不写入回放。

        业务组和 states 的 group_id 对应；离去沿用到达时的三芯成员。
        每次事件核对回放与真实资源/功率，避免在导出阶段凭空复制三芯占用。
        """
        bid = int(event.m_id)
        arrival = bool(event.m_eventType['Arrival'])
        # 仅导出真实经过观测链路的已接入业务；全网到达/阻塞计数保留在 metrics。
        if arrival:
            if blocked or not any((a, b) in self.link_index for a, b in zip(event.m_workPath, event.m_workPath[1:])):
                return
        elif bid not in self.active:
            return
        if arrival:
            service = dict(group_id=bid, source=int(event.m_sourceNode), destination=int(event.m_destNode),
                direction='forward' if event.m_sourceNode < event.m_destNode else 'backward',
                arrival_time=float(event.m_time), holding_time=float(event.m_holdTime),
                scheduled_end_time=float(event.m_time + event.m_holdTime),
                status='blocked' if blocked else 'accepted', allocation=None,
                release_time=None, release_reason=None)
            self.business_groups.append(service)
            if not blocked:
                path = [int(n) for n in event.m_workPath]
                frequency = float(self.sim.available_channel[event.m_ocuppiedwave])
                service['allocation'] = dict(path=path,
                    hops=[dict(link=[a, b], direction='forward' if a < b else 'backward',
                               cores=[int(c) for c in cores],
                               physical_cores=[CORE_TO_PHYSICAL[c] for c in cores])
                          for a, b, cores in zip(path, path[1:], event.m_ocuppiedcore)
                          if (a, b) in self.link_index],
                    channel_index=int(event.m_ocuppiedwave - self.sim.quantum_wave_num),
                    frequency_hz=frequency, wavelength_nm=299792458 / frequency * 1e9)
                self.active[bid] = service
        else:
            service = self.active.pop(bid)
            service.update(release_time=float(event.m_time), release_reason='natural')
        if not blocked:
            kind = 'arrival' if arrival else 'leave'
            self._update(service, kind)
            self.states.append(dict(time=float(event.m_time), event=kind,
                                    group_id=bid, occupancy=deepcopy(self.occupancy)))
        self._check_state()

    def _update(self, service, kind):
        """回放三芯归属一起更新；到达须空闲，离去须属于同一组，否则报错。"""
        allocation = service['allocation']
        wave = allocation['channel_index']
        for hop in allocation['hops']:
            a, b = hop['link']
            for core in hop['cores']:
                row = self.occupancy[self.link_index[a, b]][core]
                expected = -1 if kind == 'arrival' else service['group_id']
                if row[wave] != expected:
                    raise ValueError('Trace ownership disagrees with simulation lifecycle')
                row[wave] = service['group_id'] if kind == 'arrival' else -1

    def _check_state(self):
        """只读核查两方向的三芯一致性及回放占用；不会修复或改变仿真状态。"""
        sim = self.sim
        for index, (a, b) in enumerate(self.links):
            resources = sim.m_resourceMap[a, b, :, sim.quantum_wave_num:]
            powers = sim.P_link[a, b, :, sim.quantum_wave_num:]
            recorded = np.asarray(self.occupancy[index])
            actual = np.where(recorded >= 0, 2, np.where(recorded == -1, 1, 0))
            cores = sim.classical_forward_cores if a < b else sim.classical_backward_cores
            if (not np.array_equal(resources, actual)
                    or not np.all(recorded[cores] == recorded[cores[0]])
                    or not np.allclose(powers, np.where(resources == 2, sim.launch_power, 0),
                                       rtol=1e-6, atol=0)):
                raise ValueError('Three-core trace disagrees with actual resource or per-core power')

    def finish(self, metrics, source_hashes):
        """返回回放；终点仍活跃组标记 horizon 释放，不修改仿真资源和统计。

        scheduled_end_time 保留自然结束计划，release_time 是回放中的释放时间。
        时间为仿真单位，映射实际秒数由使用 JSON 的程序处理。
        """
        self._check_state()
        for bid in sorted(self.active):
            self.active[bid].update(release_time=float(self.sim.Ts), release_reason='horizon')
            self._update(self.active[bid], 'leave')
        if self.active:
            self.states.append(dict(time=float(self.sim.Ts), event='horizon_release',
                                    group_ids=sorted(self.active), occupancy=deepcopy(self.occupancy)))
        self.active.clear()
        return dict(schema_version=TRACE_SCHEMA_VERSION, kind='qkd_resource_timeline',
                    config=self.config, source_sha256=source_hashes,
                    traffic_sha256=metrics['traffic_sha256'], metrics=metrics,
                    business_groups=self.business_groups,
                    initial_occupancy=self.initial, states=self.states)


def write_trace(path, data):
    """以 UTF-8 写新 JSON，states 每个状态占一行，返回文件 SHA-256。

    使用排他新建模式防止覆盖已有结果，不是文件系统只读保护。
    拒绝 NaN/Infinity，以便其他程序读取；manifest 等字典也可使用此函数。
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf-8') as stream:
        stream.write('{\n')
        for index, (key, value) in enumerate(data.items()):
            if index:
                stream.write(',\n')
            stream.write(json.dumps(key) + ':')
            if key == 'states':
                stream.write('[\n')
                for i, row in enumerate(value):
                    if i:
                        stream.write(',\n')
                    stream.write(json.dumps(row, ensure_ascii=False, separators=(',', ':'), allow_nan=False))
                stream.write('\n]')
            else:
                stream.write(json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False))
        stream.write('\n}\n')
    return sha256(path.read_bytes()).hexdigest()


def export_excel(path, tables):
    """写普通黑白汇总表：条件、算法、四项指标及四项相对 first-fit 的比值。

    tables 为 (表名, 已汇总记录, 条件字段) 列表；条件字段是 (键, 显示名) 序列，
    同表内用全部条件配对 FF。记录使用普通扫描的指标键，SKR 输入 bit/s、输出
    kbit/s；OSNR 数值列为 dB，比值使用跨种子平均线性 OSNR，不能相除 dB。
    比值为算法均值/FF均值，不减1；分子/基准缺失或基准为零时留空。
    FF 协同度恒为零，因此协同度比值列留空。只写汇总，不写配置、样本或图表。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    metrics = [('skr_mean', 'SKR (kbit/s)', .001, '0.000'),
               ('osnr_db_mean', 'OSNR (dB)', 1, '0.000'),
               ('blocking_rate', 'Blocking rate', 1, '0.00%'),
               ('synergy_vs_FF', 'Synergy', 1, '0.000000')]
    ratio_keys = ('skr_mean', 'osnr_linear_mean', 'blocking_rate', 'synergy_vs_FF')
    for name, records, conditions in tables:
        if len(records) > 1_048_575:
            raise ValueError(f'{name} exceeds Excel row limit')
        sheet = workbook.create_sheet(name)
        headers = [label for _, label in conditions] + ['Algorithm']
        headers += [label for _, label, _, _ in metrics]
        headers += ['SKR / FF', 'OSNR / FF (linear)', 'Blocking / FF', 'Synergy / FF']
        sheet.append(headers)
        baseline = {tuple(row[key] for key, _ in conditions): row for row in records
                    if row['algorithm'] in ('FF', 'first-fit')}
        for row in records:
            key = tuple(row[key] for key, _ in conditions)
            ref = baseline.get(key, {})
            values = list(key) + [row['algorithm']]
            values += [row[metric] * factor if row[metric] is not None else None
                       for metric, _, factor, _ in metrics]
            values += [row[metric] / ref[metric]
                       if row.get(metric) is not None and ref.get(metric) not in (None, 0)
                       else None for metric in ratio_keys]
            sheet.append(values)
        formats = ['General'] * len(conditions) + ['General']
        formats += [fmt for _, _, _, fmt in metrics] + ['0.000000'] * 4
        for column, (label, number_format) in enumerate(zip(headers, formats), 1):
            sheet.column_dimensions[get_column_letter(column)].width = 26 if label == 'Algorithm' else 20
            sheet.cell(1, column).alignment = Alignment(wrap_text=True, vertical='center')
            for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                for cell in cells:
                    cell.number_format = number_format
        sheet.row_dimensions[1].height = 30
    workbook.save(path)


def annotate_max_gap(ax, points, metric, unit):
    """points 为 (横轴值, 提出算法值, CCA值)，数值已转换为图中单位。

    只比较同一横轴、同一场景的有限均值。SKR 选择 |提出算法/CCA-1| 最大点，
    以无符号百分比标注，CCA<=0 时比例未定义，跳过。OSNR 选 dB 绝对差最大点，
    标签保留提出算法减CCA的正负号。同分选较小横轴，全相等标0，缺配对则不标。
    """
    pairs = [(x, proposed, cca) for x, proposed, cca in points
             if proposed is not None and cca is not None
             and np.isfinite(proposed) and np.isfinite(cca)]
    if not pairs:
        return
    if metric == 'SKR':
        pairs = [p for p in pairs if p[2] > 0]
        if not pairs:
            return
        x, proposed, cca = max(pairs, key=lambda p: (abs(p[1] / p[2] - 1), -p[0]))
        label = f'Max SKR difference vs CCA: {abs(proposed / cca - 1):.2%}'
    else:
        x, proposed, cca = max(pairs, key=lambda p: (abs(p[1] - p[2]), -p[0]))
        label = f'Max {metric} gap vs CCA: {proposed - cca:+.3g} {unit}'
    difference = proposed - cca
    if difference != 0:
        ax.annotate('', xy=(x, proposed), xytext=(x, cca),
                    arrowprops=dict(arrowstyle='<->', color='#555555', lw=1.2))
    ax.annotate(label,
                xy=(x, (proposed + cca) / 2), xytext=(.03, .97),
                textcoords='axes fraction', ha='left', va='top', fontsize=9,
                arrowprops=dict(arrowstyle='-', color='#555555', lw=.8),
                bbox=dict(facecolor='white', edgecolor='none', alpha=.85, pad=2))


def plot_scan(output, summary, scan_axis='load'):
    """负载扫描每场景一张、功率/距离扫描每组一张 2×2 SVG：协同度、OSNR、SKR 和阻塞率。

    只使用 summary 的均值和跨种子样本标准差，不从回放重算。
    横轴由 scan_axis 选择负载/Erlang、每芯每信道功率/dBm 或统一边长/km。SKR 从 bit/s 转为 kbit/s，阻塞率显示为百分比。
    标准差有有限值时绘制阴影；调用方将单种子标准差设为空，缺失值保留断点。
    功率/距离扫描由调用方保证每个横轴值只对应一个场景。
    仅 SKR、OSNR 面板标注 GREEDY_MIN_NOISE 与 CCA 的最大差：SKR 为绝对
    百分比差，OSNR 为带符号 dB 差；同分选较小横轴，缺配对不标注。返回 SVG 文件名列表。
    """
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter

    colors = {'FF': '#D55E00', 'SCWA': '#0072B2', 'GREEDY_MIN_NOISE': '#009E73',
              'CQLI': '#E69F00', 'CCA': '#CC79A7'}
    frame = pd.DataFrame(summary)
    artifacts = []
    x_key, x_label = {
        'load': ('offered_load_erlang', 'Offered traffic (Erlang)'),
        'power': ('power_dbm', 'Power per core per channel (dBm)'),
        'distance': ('length_km', 'Uniform link length (km)'),
    }[scan_axis]
    # 物理量扫描必须将不同场景的点连成一条曲线；配对统计仍在各场景内完成。
    groups = frame.groupby('scenario', sort=False) if scan_axis == 'load' else [(
        f"{scan_axis.capitalize()} sweep | A={frame.offered_load_erlang.iloc[0]:g} Erlang | " + (
            f"observed length={frame.observed_length_m.iloc[0] / 1000:g} km" if scan_axis == 'power'
            else f"power={frame.power_dbm.iloc[0]:g} dBm/core/channel"), frame)]
    for index, (scenario, group) in enumerate(groups, 1):
        fig, axes = plt.subplots(2, 2, figsize=(12, 8), layout='constrained')
        specs = [('synergy_vs_FF', 'Signed synergy vs first-fit', 1),
                 ('osnr_db_mean', 'Reference link OSNR (dB)', 1),
                 ('skr_mean', 'Mean link SKR per channel (kbit/s)', .001),
                 ('blocking_rate', 'Blocking probability', 1)]
        for ax, (metric, label, factor) in zip(axes.flat, specs):
            for algorithm, data in group.groupby('algorithm', sort=False):
                data = data.sort_values(x_key)
                x, y = data[x_key].to_numpy(), data[metric].to_numpy(dtype=float)*factor
                sd = data[metric + '_sd'].to_numpy(dtype=float)*factor
                color = colors[algorithm]
                ax.plot(x, y, marker='o', ms=4, lw=1.7, label=algorithm, color=color)
                if np.isfinite(sd).any():
                    ax.fill_between(x, y - sd, y + sd, color=color, alpha=.12)
            ax.set_ylabel(label)
        greedy = group[group.algorithm == 'GREEDY_MIN_NOISE'].set_index(x_key)
        cca = group[group.algorithm == 'CCA'].set_index(x_key)
        for ax, metric, factor, label, unit in (
                (axes[1, 0], 'skr_mean', .001, 'SKR', 'kbit/s'),
                (axes[0, 1], 'osnr_db_mean', 1, 'OSNR', 'dB')):
            paired = greedy[[metric]].join(cca[[metric]], lsuffix='_greedy', rsuffix='_cca', how='inner').astype(float)
            annotate_max_gap(ax, [(x, row[metric + '_greedy'] * factor, row[metric + '_cca'] * factor)
                                 for x, row in paired.iterrows()], label, unit)
        axes[1, 1].yaxis.set_major_formatter(PercentFormatter(1))
        axes[0, 0].axhline(0, color="#AAAAAA", lw=.8)
        axes[0, 0].legend(fontsize=8)
        for ax in axes.flat:
            ax.set_xlabel(x_label)
            ax.grid(alpha=.2)
            ax.spines[['top', 'right']].set_visible(False)
        seeds = int(group.seed_count.iloc[0])
        fig.suptitle(f'{scenario} | {seeds} seed(s) | shaded: across-seed SD', fontsize=13)
        path = output / f'traffic_scan_{index}.svg'
        fig.savefig(path)
        artifacts.append(path.name)
        plt.close(fig)
    return artifacts


def export_business_summary(output, runs, metadata, summary):
    """生成黑白 LoadSweep/PowerSweep 简表及四指标的 PNG、SVG 趋势图。

    工作簿只含条件、算法、四项指标及各自相对FF的比值，不嵌图；完整数据留在JSON。
    图中仅标注提出算法与CCA的最大SKR绝对百分比差、OSNR差（dB），
    单图和总览共用同一规则。
    标准差仍显示为误差棒；缺失均值处断线，重合曲线不平移，不添加额外差异标注。
    """
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter

    specs = [
        ('load_scan', 'LoadSweep', 'load_erlang', 'Three-core group traffic (Erlang)',
         f"Load sweep | {metadata['fixed_power_dbm']:g} dBm/core/channel"),
        ('power_scan', 'PowerSweep', 'power_dbm', 'Power per core per channel (dBm)',
         f"Power sweep | {metadata['fixed_load_erlang']:g} Erlang of three-core groups")]
    series = [('ff', 'first-fit', '2563EB', 'o', '-'),
              ('cca', 'CCA', 'CC79A7', '^', '-.'),
              ('greedy', 'greedy_min_noise', 'D97706', 's', '--')]
    series = [item for item in series if any(r['algorithm'] == ('first-fit' if item[0] == 'ff' else 'greedy_min_noise' if item[0] == 'greedy' else 'cca') for r in runs)]
    path = output / 'skr_summary.xlsx'
    tables = []
    algorithms = {'ff': 'first-fit', 'cca': 'CCA', 'greedy': 'GREEDY_MIN_NOISE'}
    for group, sheet, *_ in specs:
        records = []
        for row in summary[group]:
            for prefix, *_ in series:
                record = dict(load_erlang=row['load_erlang'], power_dbm=row['power_dbm'],
                              algorithm=algorithms[prefix],
                              skr_mean=row[prefix + '_skr_kbit_s'] * 1000
                              if row[prefix + '_skr_kbit_s'] is not None else None)
                for key in ('osnr_db_mean', 'osnr_linear_mean', 'blocking_rate', 'synergy_vs_FF'):
                    record[key] = row[prefix + '_' + key]
                records.append(record)
        tables.append((sheet, records, [('load_erlang', 'Load (Erlang)'), ('power_dbm', 'Power (dBm)')]))
    export_excel(path, tables)
    artifacts = ['skr_summary.xlsx']
    note = (f"{runs[0]['observed_length_m'] / 1000:g} km bidirectional | slots={metadata['slots']}, warmup={metadata['warmup']} | "
            f"{len(metadata['seeds'])} seed(s); " +
            ('error bars: across-seed SD' if len(metadata['seeds']) > 1 else 'single-seed trend'))
    metrics = [('osnr_db_mean', 'osnr_db_mean_sd', 'Reference link OSNR (dB)', 'osnr_trends.png', 'classical_osnr.png'),
               ('synergy_vs_FF', 'synergy_vs_FF_sd', 'Signed synergy vs first-fit', 'synergy_trends.png', 'synergy.png'),
               ('skr_kbit_s', 'skr_sd_kbit_s', 'Mean link SKR per channel (kbit/s)', 'skr_trends.png', 'total_skr.png'),
               ('blocking_rate', 'blocking_rate_sd', 'Classical blocking probability',
                'blocking_trends.png', 'blocking_rate.png')]
    for metric, sd_key, y_label, overview, filename in metrics:
        blocking = metric == 'blocking_rate'
        upper = max((r[prefix + '_' + metric] + (r[prefix + '_' + sd_key] or 0)
                     for rows in summary.values() for r in rows for prefix, *_ in series
                     if r[prefix + '_' + metric] is not None), default=0)
        lower = min((r[prefix + '_' + metric] - (r[prefix + '_' + sd_key] or 0)
                     for rows in summary.values() for r in rows for prefix, *_ in series
                     if r[prefix + '_' + metric] is not None), default=0)
        lower = min(0, lower * 1.08)
        upper = min(1, max(.01, upper) * 1.15) if blocking else max(.01, upper * 1.08)
        fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), layout='constrained')
        for ax, (group, sheet_name, x_key, x_label, title) in zip(axes, specs):
            single, single_ax = plt.subplots(figsize=(7, 4.8), layout='constrained')
            for prefix, label, color, marker, style in series:
                points = summary[group]
                if not any(r[prefix + '_' + metric] is not None for r in points):
                    continue
                x = [r[x_key] for r in points]
                y = [r[prefix + '_' + metric] if r[prefix + '_' + metric] is not None else np.nan
                     for r in points]
                sd = [r[prefix + '_' + sd_key] for r in points]
                for target in (ax, single_ax):
                    target.plot(x, y, marker=marker, ms=7 if marker == 's' else 5,
                                markerfacecolor='none' if marker == 's' else '#' + color,
                                ls=style, lw=1.8, label=label, color='#' + color)
                    if all(v is not None for v in sd):
                        target.errorbar(x, y, yerr=sd, fmt='none', capsize=3, color='#' + color)
            if blocking:
                for target in (ax, single_ax):
                    target.yaxis.set_major_formatter(PercentFormatter(1))
            elif metric in ('skr_kbit_s', 'osnr_db_mean'):
                label, unit = ('SKR', 'kbit/s') if metric == 'skr_kbit_s' else ('OSNR', 'dB')
                points = [(row[x_key], row['greedy_' + metric], row['cca_' + metric])
                          for row in summary[group]]
                for target in (ax, single_ax):
                    annotate_max_gap(target, points, label, unit)
            for target in (ax, single_ax):
                target.set(xlabel=x_label, ylabel=y_label, title=title, ylim=(lower, upper))
                target.grid(alpha=.2)
                target.spines[['top', 'right']].set_visible(False)
                target.legend(fontsize=9, loc='lower right' if blocking else 'best')
            single.suptitle(note, fontsize=8)
            figure_path = f'{group}/{filename}'
            single.savefig(output / figure_path, dpi=180)
            vector_path = str(Path(figure_path).with_suffix('.svg'))
            single.savefig(output / vector_path)
            artifacts.append(vector_path)
            plt.close(single)
            artifacts.append(figure_path)
        fig.suptitle(note, fontsize=10)
        fig.savefig(output / overview, dpi=180)
        vector_path = str(Path(overview).with_suffix('.svg'))
        fig.savefig(output / vector_path)
        artifacts.append(vector_path)
        plt.close(fig)
        artifacts.append(overview)
    return artifacts


def export_scan_results(output, data, save_samples=False):
    """输出 traffic_scan.json、xlsx 和四指标对比 SVG；JSON 始终保留 samples，不重算统计。"""
    (output / 'traffic_scan.json').write_text(
        json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    export_excel(output / 'traffic_scan.xlsx', [('Summary', data['summary'], [
        ('offered_load_erlang', 'Load (Erlang)'), ('power_dbm', 'Power (dBm)'),
        ('observed_length_m', 'Link length (m)')])])
    return plot_scan(output, data['summary'], data['config'].get('scan_axis', 'load'))
